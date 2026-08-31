# -*- coding: utf-8 -*-
"""Vue API：任务安排（看板/KPI/单任务/批量/导入导出）"""
from datetime import date, datetime, timedelta

from models import db, Customer, InspectionTask, Inspector, User


def test_task_work_calendar_uses_published_holiday_source(admin_client):
    response = admin_client.get('/api/task-schedule/work-calendar')
    assert response.status_code == 200
    data = response.get_json()['data']
    assert data['version'] == 'cn-gov-2026-v1'
    assert data['covered_years'] == [2026]
    holidays = {item['date']: item['name'] for item in data['holidays']}
    workdays = {item['date']: item['name'] for item in data['makeup_workdays']}
    assert holidays['2026-06-19'] == '端午节'
    assert workdays['2026-02-14'] == '春节调休上班'


def test_task_timing_payload_uses_workday_execution_to_approval_window():
    from services.task_schedule_service import task_timing_payload

    task = InspectionTask(
        title='计时任务', status='已完成',
        actual_start=datetime(2026, 8, 24, 8, 0),
        actual_end=datetime(2026, 8, 24, 18, 0),
    )
    timing = task_timing_payload(task)
    assert timing['actual_duration_hours'] == 7.5
    assert timing['actual_duration_text'] == '7小时30分钟'
    assert timing['actual_effort'] == 0.94


def test_task_timing_excludes_night_lunch_and_weekend():
    from services.task_schedule_service import task_timing_payload

    overnight = InspectionTask(
        title='跨夜任务', status='已完成',
        actual_start=datetime(2026, 8, 24, 9, 9),
        actual_end=datetime(2026, 8, 25, 10, 13),
    )
    timing = task_timing_payload(overnight)
    assert timing['actual_duration_text'] == '8小时34分钟'
    assert timing['actual_effort'] == 1.07

    eight_hours = InspectionTask(
        title='满一人天任务', status='已完成',
        actual_start=datetime(2026, 8, 24, 8, 30),
        actual_end=datetime(2026, 8, 25, 9, 0),
    )
    timing = task_timing_payload(eight_hours)
    assert timing['actual_duration_text'] == '8小时'
    assert timing['actual_effort'] == 1.0

    weekend = InspectionTask(
        title='跨周末任务', status='已完成',
        actual_start=datetime(2026, 8, 28, 16, 30),  # Friday
        actual_end=datetime(2026, 8, 31, 9, 30),     # Monday
    )
    timing = task_timing_payload(weekend)
    assert timing['actual_duration_text'] == '2小时'
    assert timing['actual_effort'] == 0.25


def _seed(app):
    with app.app_context():
        from services.task_schedule_service import local_now
        c = Customer(name='看板客户')
        db.session.add(c)
        db.session.flush()
        op = User.create_with_password(username='tsuser', password='x', role='operator', realname='张工')
        db.session.add(op)
        db.session.flush()
        db.session.add(Inspector(user_id=op.id, is_active=True))
        db.session.add(InspectionTask(title='2026年二季度巡检', customer_id=c.id, status='待执行',
                                      assigned_to_user_id=op.id,
                                      planned_start=date.today() - timedelta(days=10),
                                      planned_end=date.today() - timedelta(days=1)))
        db.session.add(InspectionTask(title='2026年三季度巡检', customer_id=c.id, status='执行中',
                                      assigned_to_user_id=op.id,
                                      actual_start=local_now() - timedelta(hours=2),
                                      planned_start=date.today(), planned_end=date.today() + timedelta(days=30)))
        db.session.commit()
        return c.id, op.id


class TestTaskScheduleApi:
    def test_board_default_quarter(self, admin_client, app, monkeypatch):
        fake_now = datetime(2026, 8, 25, 10, 30)
        monkeypatch.setattr('services.task_schedule_service.local_now', lambda: fake_now)
        _, op_id = _seed(app)
        r = admin_client.get('/api/task-schedule')
        assert r.get_json()['code'] == 0
        d = r.get_json()['data']
        # 默认工程师视角
        assert d['view'] == 'engineer'
        assert d['kpi']['total'] == 2
        assert len(d['engineer_groups'][str(op_id)]) == 2
        assert d['kpi']['overdue'] == 1
        assert d['kpi']['act_effort'] == 0.25
        assert len(d['engineers']) == 1
        running = next(t for t in d['tasks'] if t['status'] == '执行中')
        assert running['actual_duration_text'] == '2小时'
        assert running['actual_effort'] == 0.25

    def test_explicit_all_period_does_not_fall_back_to_current_quarter(
            self, admin_client, app):
        _seed(app)
        with app.app_context():
            customer_id = Customer.query.filter_by(name='看板客户').one().id
            db.session.add(InspectionTask(
                title='历史跨季度任务', customer_id=customer_id, status='已完成',
                planned_start=date(2020, 1, 1), planned_end=date(2020, 1, 2)))
            db.session.commit()

        data = admin_client.get('/api/task-schedule?period=').get_json()['data']
        assert '历史跨季度任务' in {item['title'] for item in data['tasks']}

    def test_status_view(self, admin_client, app):
        _seed(app)
        r = admin_client.get('/api/task-schedule?view=status')
        d = r.get_json()['data']
        assert d['view'] == 'status'
        assert d['status_groups']['待执行'] and d['status_groups']['执行中']
        assert '已安排' in d['status_groups']
        assert '待审核' in d['status_groups']  # V21 新状态列

    def test_scheduled_requires_task_deadline_and_does_not_start_timer(
            self, admin_client, app):
        cid, op_id = _seed(app)
        created = admin_client.post('/api/task-schedule', json={
            'title': '待排期任务', 'customer_id': cid,
        })
        tid = created.get_json()['data']['id']

        missing_period = admin_client.put(
            f'/api/task-schedule/{tid}', json={'status': '已安排'})
        assert missing_period.status_code == 400
        assert '完整的任务期限' in missing_period.get_json()['message']

        missing_assignee = admin_client.put(f'/api/task-schedule/{tid}', json={
            'status': '已安排',
            'scheduled_start': '2026-08-25',
            'scheduled_end': '2026-08-29',
        })
        assert missing_assignee.status_code == 400
        assert '负责人' in missing_assignee.get_json()['message']

        scheduled = admin_client.put(f'/api/task-schedule/{tid}', json={
            'status': '已安排',
            'assignee_id': op_id,
            'planned_start': '2026-07-01',
            'planned_end': '2026-09-30',
            'scheduled_start': '2026-08-25',
            'scheduled_end': '2026-08-29',
        })
        assert scheduled.status_code == 200
        with app.app_context():
            task = db.session.get(InspectionTask, tid)
            assert task.status == '已安排'
            assert task.planned_start == date(2026, 7, 1)
            assert task.planned_end == date(2026, 9, 30)
            assert task.scheduled_start == date(2026, 8, 25)
            assert task.scheduled_end == date(2026, 8, 29)
            assert task.actual_start is None
            assert task.actual_end is None
            assert task.actual_effort is None
        clear_assignee = admin_client.put(
            f'/api/task-schedule/{tid}', json={'assignee_id': None})
        assert clear_assignee.status_code == 400
        assert '负责人' in clear_assignee.get_json()['message']

        board = admin_client.get('/api/task-schedule?period=&view=status').get_json()['data']
        assert board['kpi']['scheduled'] == 1
        item = board['status_groups']['已安排'][0]
        assert item['id'] == tid
        assert item['planned_start'] == '2026-07-01'
        assert item['planned_end'] == '2026-09-30'
        assert item['scheduled_start'] == '2026-08-25'
        assert item['scheduled_end'] == '2026-08-29'

        assert admin_client.put(
            f'/api/task-schedule/{tid}', json={'status': '执行中'}).status_code == 200
        with app.app_context():
            task = db.session.get(InspectionTask, tid)
            assert task.actual_start is not None

        assert admin_client.put(
            f'/api/task-schedule/{tid}', json={'status': '已安排'}).status_code == 200
        with app.app_context():
            task = db.session.get(InspectionTask, tid)
            assert task.actual_start is None
            assert task.actual_end is None
            assert task.actual_effort is None

    def test_reviewing_kpi_and_group(self, admin_client, app):
        """V21: 待审核任务计入 KPI + 状态分组 + 排序优先级"""
        cid, op_id = _seed(app)
        with app.app_context():
            from datetime import date as _d
            db.session.add(InspectionTask(title='待审核任务', customer_id=cid, status='待审核',
                                          assigned_to_user_id=op_id,
                                          planned_start=_d.today(),
                                          planned_end=_d.today() + __import__('datetime').timedelta(days=5)))
            db.session.commit()
        r = admin_client.get('/api/task-schedule?view=status')
        d = r.get_json()['data']
        assert d['kpi']['reviewing'] == 1
        assert len(d['status_groups']['待审核']) == 1
        # 排序：逾期(待执行) → 执行中 → 待审核（同级按截止时间升序）
        titles = [t['title'] for t in d['tasks']]
        assert titles.index('2026年二季度巡检') < titles.index('2026年三季度巡检') \
            < titles.index('待审核任务')

    def test_contract_review_flow_and_notification(self, admin_client, op_client, app, monkeypatch):
        cid, _ = _seed(app)
        with app.app_context():
            customer = db.session.get(Customer, cid)
            customer.contract_end_date = date.today() - timedelta(days=1)
            db.session.commit()
        events = []
        monkeypatch.setattr(
            'utils.wecom_notify.wecom_broadcast',
            lambda event_type, *args, **kwargs: events.append(event_type) or (1, 0))
        created = op_client.post('/api/task-schedule', json={
            'title': '过期合同巡检', 'customer_id': cid,
            'contract_exception_reason': '紧急安全检查',
            'planned_start': date.today().isoformat(),
            'planned_end': (date.today() + timedelta(days=7)).isoformat(),
        })
        assert created.status_code == 200
        task_id = created.get_json()['data']['id']
        board = admin_client.get('/api/task-schedule?view=status').get_json()['data']
        assert board['kpi']['contract_review'] == 1
        assert board['status_groups']['合同审批'][0]['contract_exception_reason'] == '紧急安全检查'
        assert 'contract_review' in events

        # 通用更新接口不能绕过专用审核入口。
        bypass = admin_client.put(
            f'/api/task-schedule/{task_id}', json={'status': '待执行'})
        assert bypass.status_code == 400
        assert '审核接口' in bypass.get_json()['message']
        assert op_client.post(
            f'/api/task-schedule/{task_id}/contract-review',
            json={'approved': True}).status_code == 403

        reviewed = admin_client.post(
            f'/api/task-schedule/{task_id}/contract-review',
            json={'approved': True, 'comment': '同意本次例外'})
        assert reviewed.status_code == 200
        assert reviewed.get_json()['data']['status'] == '待执行'
        with app.app_context():
            from models import AuditLog, Notification
            task = db.session.get(InspectionTask, task_id)
            assert task.contract_exception_status == '通过'
            assert '审核意见：同意本次例外' in task.contract_exception_reason
            assert AuditLog.query.filter_by(
                action='task:contract_review', target_id=task_id).count() == 1
            assert Notification.query.filter(
                Notification.title.contains('任务合同例外审核通过')).count() >= 1

    def test_contract_review_rejects_to_cancelled(self, admin_client, app):
        cid, _ = _seed(app)
        with app.app_context():
            task = InspectionTask(
                title='拒绝例外', customer_id=cid, status='合同审批',
                contract_exception_status='待审核', contract_exception_reason='无合同',
                contract_exception_by='op')
            db.session.add(task)
            db.session.commit()
            task_id = task.id
        r = admin_client.post(
            f'/api/task-schedule/{task_id}/contract-review',
            json={'approved': False})
        assert r.status_code == 200
        assert r.get_json()['data']['status'] == '已取消'

    def test_contract_review_approved_keeps_complete_task_deadline(
            self, admin_client, app):
        cid, op_id = _seed(app)
        with app.app_context():
            task = InspectionTask(
                title='已有任务期限的例外任务', customer_id=cid, status='合同审批',
                assigned_to_user_id=op_id,
                planned_start=date(2026, 7, 1), planned_end=date(2026, 9, 30),
                scheduled_start=date(2026, 8, 25), scheduled_end=date(2026, 8, 29),
                contract_exception_status='待审核', contract_exception_reason='临时例外',
                contract_exception_by='op')
            db.session.add(task)
            db.session.commit()
            task_id = task.id

        response = admin_client.post(
            f'/api/task-schedule/{task_id}/contract-review',
            json={'approved': True})
        assert response.status_code == 200
        assert response.get_json()['data']['status'] == '已安排'
        with app.app_context():
            task = db.session.get(InspectionTask, task_id)
            assert task.actual_start is None

    def test_status_machine_validation(self, admin_client, app):
        """V21: 状态机校验 — 已取消不可回退；待审核不可手工重复"""
        cid, _ = _seed(app)
        with app.app_context():
            t = InspectionTask(title='取消任务', customer_id=cid, status='已取消')
            db.session.add(t)
            db.session.commit()
            tid = t.id
        r = admin_client.put(f'/api/task-schedule/{tid}', json={'status': '执行中'})
        assert r.status_code == 400
        assert '不允许' in r.get_json()['message']

    def test_board_sort_order(self, admin_client, app):
        """看板排序：逾期最前 → 执行中 → 待执行 → 已完成。"""
        cid, op_id = _seed(app)
        with app.app_context():
            t_overdue_running = InspectionTask(
                title='逾期执行中', customer_id=cid, status='执行中',
                assigned_to_user_id=op_id,
                planned_start=date.today() - timedelta(days=8),
                planned_end=date.today() - timedelta(days=1))
            t_done = InspectionTask(
                title='已完成', customer_id=cid, status='已完成',
                assigned_to_user_id=op_id,
                planned_start=date.today(), planned_end=date.today() + timedelta(days=5))
            t_pending_future = InspectionTask(
                title='未来待执行', customer_id=cid, status='待执行',
                assigned_to_user_id=op_id,
                planned_start=date.today() + timedelta(days=3),
                planned_end=date.today() + timedelta(days=10))
            db.session.add_all([t_overdue_running, t_done, t_pending_future])
            db.session.commit()
        r = admin_client.get('/api/task-schedule')
        d = r.get_json()['data']
        titles = [t['title'] for t in d['tasks']]
        # 逾期执行中 → 逾期待执行（seed 任务1） → 执行中（seed 任务2） → 未来待执行 → 已完成
        assert titles == ['逾期执行中', '2026年二季度巡检', '2026年三季度巡检',
                          '未来待执行', '已完成']
        # 工程师视角组内顺序一致
        r = admin_client.get('/api/task-schedule?view=engineer')
        d = r.get_json()['data']
        assert [t['title'] for t in d['engineer_groups'][str(op_id)]] == titles

    def test_board_uses_stage_specific_start_time_within_each_status(
            self, admin_client, app):
        cid, op_id = _seed(app)
        future_end = date.today() + timedelta(days=180)
        with app.app_context():
            tasks = [
                InspectionTask(
                    title='排期后', customer_id=cid, status='已安排',
                    assigned_to_user_id=op_id,
                    planned_start=date(2026, 7, 1), planned_end=future_end,
                    scheduled_start=date(2026, 9, 7), scheduled_end=date(2026, 9, 11)),
                InspectionTask(
                    title='排期前', customer_id=cid, status='已安排',
                    assigned_to_user_id=op_id,
                    planned_start=date(2026, 8, 1), planned_end=future_end,
                    scheduled_start=date(2026, 8, 31), scheduled_end=date(2026, 9, 4)),
                InspectionTask(
                    title='排期缺失', customer_id=cid, status='已安排',
                    assigned_to_user_id=op_id,
                    planned_start=date(2026, 6, 1), planned_end=future_end),
                InspectionTask(
                    title='合同后', customer_id=cid, status='待执行',
                    assigned_to_user_id=op_id,
                    planned_start=date(2026, 8, 1), planned_end=future_end),
                InspectionTask(
                    title='合同前', customer_id=cid, status='待执行',
                    assigned_to_user_id=op_id,
                    planned_start=date(2026, 7, 1), planned_end=future_end),
                InspectionTask(
                    title='实施后', customer_id=cid, status='已完成',
                    assigned_to_user_id=op_id,
                    actual_start=datetime(2026, 9, 7, 8, 30),
                    actual_end=datetime(2026, 9, 7, 17, 30)),
                InspectionTask(
                    title='实施前', customer_id=cid, status='已完成',
                    assigned_to_user_id=op_id,
                    actual_start=datetime(2026, 8, 31, 8, 30),
                    actual_end=datetime(2026, 8, 31, 17, 30)),
            ]
            db.session.add_all(tasks)
            db.session.commit()

        groups = admin_client.get(
            '/api/task-schedule?period=&view=status').get_json()['data']['status_groups']
        assert [item['title'] for item in groups['已安排'] if item['title'].startswith('排期')] == [
            '排期前', '排期后', '排期缺失']
        assert [item['title'] for item in groups['待执行'] if item['title'].startswith('合同')] == [
            '合同前', '合同后']
        assert [item['title'] for item in groups['已完成'] if item['title'].startswith('实施')] == [
            '实施前', '实施后']

    def test_engineer_view(self, admin_client, app):
        _, op_id = _seed(app)
        r = admin_client.get('/api/task-schedule?view=engineer')
        d = r.get_json()['data']
        assert d['view'] == 'engineer'
        assert len(d['engineer_groups'][str(op_id)]) == 2

    def test_quick_add_update_delete(self, admin_client, app):
        cid, op_id = _seed(app)
        r = admin_client.post('/api/task-schedule', json={
            'title': '新增任务X', 'customer_id': cid, 'assignee_id': op_id,
            'planned_start': '2026-08-01', 'planned_end': '2026-08-31', 'priority': '高',
            'estimated_effort': 1.5, 'task_type': '计划',
            'visit_at': '2026-08-31T08:55',
        })
        assert r.get_json()['code'] == 0
        tid = r.get_json()['data']['id']
        r = admin_client.put(f'/api/task-schedule/{tid}', json={'status': '执行中'})
        assert r.get_json()['code'] == 0
        manual = admin_client.put(
            f'/api/task-schedule/{tid}', json={'actual_effort': 1})
        assert manual.status_code == 400
        assert '自动计算' in manual.get_json()['message']
        with app.app_context():
            t = db.session.get(InspectionTask, tid)
            assert t.status == '执行中'
            assert t.actual_start is not None
            assert t.actual_effort is None
            assert t.actual_end is None
            assert t.visit_at.strftime('%Y-%m-%d %H:%M') == '2026-08-31 00:55'
            t.actual_start = datetime(2026, 8, 24, 9, 0)
            t.actual_end = datetime(2026, 8, 24, 10, 0)
            t.actual_effort = 0.13
            db.session.commit()
        r = admin_client.put(f'/api/task-schedule/{tid}', json={'status': '待执行'})
        assert r.get_json()['code'] == 0
        with app.app_context():
            t = db.session.get(InspectionTask, tid)
            assert t.status == '待执行'
            assert t.actual_start is None
            assert t.actual_end is None
            assert t.actual_effort is None
        r = admin_client.put(f'/api/task-schedule/{tid}', json={'status': '执行中'})
        assert r.get_json()['code'] == 0
        with app.app_context():
            t = db.session.get(InspectionTask, tid)
            assert t.actual_start is not None
            assert t.actual_start != datetime(2026, 8, 24, 9, 0)
        r = admin_client.delete(f'/api/task-schedule/{tid}')
        assert r.get_json()['code'] == 0
        with app.app_context():
            assert db.session.get(InspectionTask, tid) is None

    def test_quick_add_requires_title_customer(self, admin_client, app):
        r = admin_client.post('/api/task-schedule', json={'title': '', 'customer_id': 1})
        assert r.status_code == 400
        r = admin_client.post('/api/task-schedule', json={'title': 'x'})
        assert r.status_code == 400

    def test_batch_actions(self, admin_client, app):
        cid, op_id = _seed(app)
        with app.app_context():
            t1 = InspectionTask(title='批量1', customer_id=cid, status='待执行')
            t2 = InspectionTask(title='批量2', customer_id=cid, status='待执行')
            db.session.add_all([t1, t2])
            db.session.commit()
            ids = [t1.id, t2.id]
        r = admin_client.post('/api/task-schedule/batch', json={'ids': ids, 'action': 'assign', 'value': op_id})
        assert r.get_json()['code'] == 0
        r = admin_client.post('/api/task-schedule/batch', json={'ids': ids, 'action': 'status', 'value': '执行中'})
        assert r.get_json()['code'] == 0
        with app.app_context():
            tasks = InspectionTask.query.filter(InspectionTask.id.in_(ids)).all()
            assert all(t.status == '执行中' and t.assigned_to_user_id == op_id for t in tasks)
        r = admin_client.post('/api/task-schedule/batch', json={
            'ids': ids, 'action': 'status', 'value': '待执行'})
        assert r.get_json()['code'] == 0
        with app.app_context():
            tasks = InspectionTask.query.filter(InspectionTask.id.in_(ids)).all()
            assert all(t.status == '待执行' and t.actual_start is None for t in tasks)
        r = admin_client.post('/api/task-schedule/batch', json={'ids': ids, 'action': 'delete'})
        assert r.get_json()['code'] == 0
        with app.app_context():
            assert InspectionTask.query.filter(InspectionTask.id.in_(ids)).count() == 0

    def test_import_template_and_import(self, admin_client, app):
        import base64
        import io

        from openpyxl import load_workbook

        cid, op_id = _seed(app)
        with app.app_context():
            op = db.session.get(User, op_id)
            op_realname = op.realname
        r = admin_client.get('/api/task-schedule/import-template')
        assert r.get_json()['code'] == 0
        assert r.get_json()['data']['filename'].endswith('.xlsx')
        template = load_workbook(io.BytesIO(base64.b64decode(r.get_json()['data']['content'])))
        template_status = template.active.cell(row=2, column=8)
        assert template_status.value == '已完成'
        assert template_status.fill.fgColor.rgb.endswith('E1F3D8')
        # 构造导入 xlsx
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(['客户名称', '任务描述', '优先级', '开始日期', '完成日期', '完成状态', '负责人', '完成时间', '预估工作量', '实际工作量'])
        ws.append(['看板客户', '导入任务A', '中', '2026-08-01', '2026-08-31', '已完成', op_realname, '', '1', '0.5'])
        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)
        r = admin_client.post('/api/task-schedule/import', data={
            'importFile': (bio, 'import.xlsx'),
        }, content_type='multipart/form-data')
        assert r.get_json()['code'] == 0
        assert '新增 1' in r.get_json()['data']['message']
        with app.app_context():
            task = InspectionTask.query.filter_by(title='导入任务A').one()
            assert task.planned_start == date(2026, 8, 1)
            assert task.planned_end == date(2026, 8, 31)
            assert task.scheduled_start is None
            assert task.scheduled_end is None

    def test_import_keeps_contract_task_and_execution_periods_separate(
            self, admin_client, app):
        import io
        from openpyxl import Workbook

        cid, op_id = _seed(app)
        with app.app_context():
            op_realname = db.session.get(User, op_id).realname
        wb = Workbook()
        ws = wb.active
        ws.append([
            '客户名称', '任务描述', '完成状态', '负责人',
            '合同时效开始日期', '合同时效结束日期',
            '任务期限开始日期', '任务期限结束日期',
        ])
        ws.append([
            '看板客户', '三时效导入任务', '已安排', op_realname,
            '2026-07-01', '2026-09-30', '2026-08-25', '2026-08-29',
        ])
        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)

        response = admin_client.post('/api/task-schedule/import', data={
            'importFile': (bio, 'three-periods.xlsx'),
        }, content_type='multipart/form-data')
        assert response.status_code == 200, response.get_json()
        with app.app_context():
            task = InspectionTask.query.filter_by(title='三时效导入任务').one()
            assert task.status == '已安排'
            assert task.planned_start == date(2026, 7, 1)
            assert task.planned_end == date(2026, 9, 30)
            assert task.scheduled_start == date(2026, 8, 25)
            assert task.scheduled_end == date(2026, 8, 29)
            assert task.actual_start is None

    def test_export_excel_supports_task_deadline_overlap_and_multiple_statuses(
            self, admin_client, app, monkeypatch):
        import base64
        import io

        from openpyxl import load_workbook

        fake_now = datetime(2026, 8, 25, 10, 30)
        monkeypatch.setattr('services.task_schedule_service.local_now', lambda: fake_now)
        customer_id, op_id = _seed(app)
        with app.app_context():
            from services.task_schedule_service import local_now
            running = InspectionTask.query.filter_by(title='2026年三季度巡检').one()
            running.planned_start = date(2026, 7, 1)
            running.planned_end = date(2026, 9, 30)
            running.scheduled_start = date(2026, 8, 30)
            running.scheduled_end = date(2026, 9, 2)
            running.actual_start = local_now() - timedelta(hours=2)

            scheduled = InspectionTask.query.filter_by(title='2026年二季度巡检').one()
            scheduled.status = '已安排'
            scheduled.planned_start = date(2026, 8, 24)
            scheduled.planned_end = date(2026, 8, 30)
            scheduled.scheduled_start = date(2026, 8, 18)
            scheduled.scheduled_end = date(2026, 8, 24)

            # 合同时效落在导出范围、但任务期限不相交，必须排除。
            db.session.add(InspectionTask(
                title='仅合同日期命中', customer_id=customer_id, status='已安排',
                assigned_to_user_id=op_id,
                planned_start=date(2026, 8, 24), planned_end=date(2026, 8, 30),
                scheduled_start=date(2026, 9, 7), scheduled_end=date(2026, 9, 11)))
            # 任务期限命中、但状态未被选择，也必须排除。
            db.session.add(InspectionTask(
                title='状态未选择', customer_id=customer_id, status='已完成',
                assigned_to_user_id=op_id,
                planned_start=date(2026, 7, 1), planned_end=date(2026, 9, 30),
                scheduled_start=date(2026, 8, 24), scheduled_end=date(2026, 8, 30)))
            db.session.commit()
        response = admin_client.get('/api/task-schedule/export', query_string={
            'period': '',
            'status': '已完成',  # 多选 statuses 应覆盖看板遗留单状态。
            'statuses': '已安排,执行中',
            'scheduled_from': '2026-08-24',
            'scheduled_to': '2026-08-30',
        })
        assert response.status_code == 200, response.get_json()
        data = response.get_json()['data']
        assert data['count'] == 2
        assert '任务期限_2026-08-24_至_2026-08-30' in data['filename']

        workbook = load_workbook(io.BytesIO(base64.b64decode(data['content'])))
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        assert rows[0] == tuple([
            '客户名称', '任务描述', '优先级',
            '合同时效开始日期', '合同时效结束日期',
            '任务期限开始日期', '任务期限结束日期',
            '完成状态', '负责人', '实施开始时间', '实施结束时间', '实施耗时',
            '预估人天', '实际人天',
        ])
        assert len(rows) == 3
        exported = {row[1]: row for row in rows[1:]}
        assert set(exported) == {'2026年二季度巡检', '2026年三季度巡检'}
        running_row = exported['2026年三季度巡检']
        assert running_row[9]
        assert running_row[10] is None
        assert running_row[11] == '2小时'
        assert running_row[13] == '0.25'
        status_cells = {sheet.cell(row=index, column=2).value:
                        sheet.cell(row=index, column=8)
                        for index in range(2, sheet.max_row + 1)}
        assert status_cells['2026年二季度巡检'].value == '已安排'
        assert status_cells['2026年二季度巡检'].fill.fgColor.rgb.endswith('D5F5F6')
        assert status_cells['2026年三季度巡检'].value == '执行中'
        assert status_cells['2026年三季度巡检'].fill.fgColor.rgb.endswith('D9ECFF')
        assert status_cells['2026年三季度巡检'].font.color.rgb.endswith('409EFF')

        with app.app_context():
            from models import AuditLog
            audit = AuditLog.query.filter_by(action='task:export').one()
            assert '任务期限 2026-08-24 至 2026-08-30' in audit.detail
            assert '状态 已安排、执行中' in audit.detail

    def test_export_rejects_reversed_or_invalid_date_range(self, admin_client):
        reversed_range = admin_client.get(
            '/api/task-schedule/export?period=&scheduled_from=2026-08-31&scheduled_to=2026-08-01')
        assert reversed_range.status_code == 400
        assert '开始日期' in reversed_range.get_json()['message']

        invalid = admin_client.get(
            '/api/task-schedule/export?period=&scheduled_from=2026-99-01')
        assert invalid.status_code == 400
        assert '格式' in invalid.get_json()['message']

        invalid_status = admin_client.get(
            '/api/task-schedule/export?period=&statuses=不存在')
        assert invalid_status.status_code == 400
        assert '任务状态无效' in invalid_status.get_json()['message']

    def test_scheduled_status_has_matching_excel_color(self, admin_client, app):
        import base64
        import io

        from openpyxl import load_workbook

        cid, op_id = _seed(app)
        response = admin_client.post('/api/task-schedule', json={
            'title': '已排期导出', 'customer_id': cid, 'assignee_id': op_id,
            'planned_start': '2026-07-01', 'planned_end': '2026-09-30',
            'scheduled_start': '2026-08-25', 'scheduled_end': '2026-08-29',
        })
        assert response.status_code == 200
        with app.app_context():
            task = db.session.get(InspectionTask, response.get_json()['data']['id'])
            assert task.status == '已安排'
            assert task.actual_start is None

        exported = admin_client.get(
            '/api/task-schedule/export?period=&statuses=已安排&scheduled_from=2026-08-25&scheduled_to=2026-08-25')
        workbook = load_workbook(io.BytesIO(base64.b64decode(
            exported.get_json()['data']['content'])))
        status_cell = workbook.active.cell(row=2, column=8)
        assert status_cell.value == '已安排'
        assert status_cell.fill.fgColor.rgb.endswith('D5F5F6')
        assert status_cell.font.color.rgb.endswith('08979C')

    def test_import_missing_required_column(self, admin_client, app):
        import io
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(['任务描述'])
        ws.append(['x'])
        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)
        r = admin_client.post('/api/task-schedule/import', data={
            'importFile': (bio, 'bad.xlsx'),
        }, content_type='multipart/form-data')
        assert r.status_code == 400

    def test_permissions(self, viewer_client, op_client):
        assert viewer_client.get('/api/task-schedule').status_code == 403  # 无 task:schedule
        assert op_client.get('/api/task-schedule').status_code == 200
        assert viewer_client.post('/api/task-schedule', json={}).status_code == 403
        assert viewer_client.get('/api/task-schedule/export?period=').status_code == 403

    def test_legacy_status_form_requires_task_permission(self, viewer_client, app):
        cid, _ = _seed(app)
        with app.app_context():
            task_id = InspectionTask.query.filter_by(customer_id=cid).first().id
        r = viewer_client.post(
            f'/task-schedule/{task_id}/status-form', json={'status': '执行中'})
        assert r.status_code == 403
        assert r.get_json()['required'] == 'task:schedule'
