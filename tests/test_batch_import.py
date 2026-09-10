# -*- coding: utf-8 -*-
"""批量导入端点：备件档案/库存/巡检记录/故障记录（模板列见 views/system.download_template）"""
import io
from datetime import date, datetime

import pytest

from domain_metadata import get_entity_schema
from models import (db, Customer, CustomerCategory, Device, Fault, Inspection,
                    NetworkType, Rack, RackInstall, Region, SparePart, SpareStock)
from utils.crypto import decrypt_password
from utils.import_templates import IMPORT_TEMPLATES, get_import_field_mapping
from utils.json_fields import parse_json


def _xlsx(headers, rows):
    """构造 xlsx bytes：headers 为列名列表，rows 为单元格列表列表"""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _template_row(module, values):
    return [values.get(field, '') for _header, field in IMPORT_TEMPLATES[module]['fields']]


class TestImportTemplates:
    def test_device_download_omits_certificate_expiry_without_shifting_examples(self, admin_client):
        import openpyxl
        response = admin_client.get('/exports/download-template/device')
        assert response.status_code == 200
        workbook = openpyxl.load_workbook(io.BytesIO(response.data))
        headers = [cell.value for cell in workbook.active[1]]
        assert '证书到期日期' not in headers
        assert '证书到期' not in headers
        for sheet in workbook:
            assert all('证书到期' not in str(cell.value or '')
                       for row in sheet for cell in row)
        examples = dict(zip(headers, [cell.value for cell in workbook.active[2]]))
        assert examples['是否维修'] == '是'  # 当前布尔下拉首项，不得残留日期值。
        assert examples['是否在用'] == '是'
        assert headers[-1] == '设备ID（可选）'
        workbook.close()

    @pytest.mark.parametrize('enforce,linked', [(False, False), (True, True), (True, False)])
    def test_device_customer_candidates_follow_import_scope(
            self, app, op_client, monkeypatch, enforce, linked):
        import openpyxl
        from models import User
        monkeypatch.setitem(app.config, 'CUSTOMER_SCOPE_ENFORCE', enforce)
        with app.app_context():
            allowed = Customer(name='设备可导入客户')
            other = Customer(name='其他客户')
            db.session.add_all([allowed, other])
            user = User.query.filter_by(username='op').one()
            user.customers = [allowed] if linked else []
            db.session.commit()
        response = op_client.get('/exports/download-template/device')
        assert response.status_code == 200
        workbook = openpyxl.load_workbook(io.BytesIO(response.data))
        sheet = workbook.active
        col = IMPORT_TEMPLATES['device']['headers'].index('客户') + 1
        cell = sheet.cell(2, col)
        validation = next(v for v in sheet.data_validations.dataValidation if cell.coordinate in v.sqref)
        if enforce and not linked:
            assert cell.value is None
            assert '当前没有可用于设备导入的客户' in validation.prompt
            assert validation.type is None
        else:
            assert validation.type == 'list'
            assert validation.showDropDown is False
            assert validation.showErrorMessage is False
            name, span = next(workbook.defined_names[validation.formula1].destinations)
            values = {c.value for row in workbook[name][span] for c in row}
            assert values == ({'设备可导入客户'} if enforce else {'设备可导入客户', '其他客户'})
            assert cell.value in values
        workbook.close()

    def test_device_customer_dropdown_uses_named_range(self, admin_client, app):
        import openpyxl
        with app.app_context():
            db.session.add(Customer(name='下拉客户完整名称'))
            db.session.commit()
        response = admin_client.get('/exports/download-template/device')
        workbook = openpyxl.load_workbook(io.BytesIO(response.data))
        sheet = workbook.active
        col = IMPORT_TEMPLATES['device']['headers'].index('客户') + 1
        validation = next(v for v in sheet.data_validations.dataValidation
                          if sheet.cell(2, col).coordinate in v.sqref)
        assert validation.type == 'list'
        assert validation.showDropDown is False
        assert validation.showErrorMessage is False
        assert '手动填写完整客户名称' in validation.prompt
        assert '不会随设备导入自动新增' in validation.prompt
        assert f'{sheet.cell(5000, col).coordinate}' in validation.sqref
        destinations = list(workbook.defined_names[validation.formula1].destinations)
        assert len(destinations) == 1
        name, cell_range = destinations[0]
        assert '下拉客户完整名称' in [cell.value for row in workbook[name][cell_range] for cell in row]
        assert workbook[name].sheet_state == 'hidden'
        # A typed value outside the download-time list survives save/reopen.
        sheet.cell(3, col, '下载后新增的客户')
        out = io.BytesIO()
        workbook.save(out)
        reopened = openpyxl.load_workbook(io.BytesIO(out.getvalue()))
        assert reopened.active.cell(3, col).value == '下载后新增的客户'
        reopened.close()
        workbook.close()

    @pytest.mark.parametrize('exists', [True, False])
    def test_device_typed_customer_is_checked_at_import(self, admin_client, app, exists):
        # Users can type a name not present when the template was downloaded.
        assert admin_client.get('/exports/download-template/device').status_code == 200
        if exists:
            with app.app_context():
                db.session.add(Customer(name='下载后新增客户'))
                db.session.commit()
        content = _xlsx(['客户', '名称'], [['下载后新增客户', '手填客户设备']])
        response = admin_client.post('/api/v2/devices/import', data={
            'import_file': (io.BytesIO(content), '手填客户设备.xlsx'),
            'mode': 'create', 'dry_run': '1',
        }, content_type='multipart/form-data')
        assert response.status_code == 200, response.get_json()
        result = response.get_json()['data']
        assert result['failed'] == (0 if exists else 1), result
        assert result['create'] == (1 if exists else 0), result
        with app.app_context():
            assert Customer.query.filter_by(name='下载后新增客户').count() == int(exists)

    def test_device_template_has_field_prompts_formats_and_validation(self, admin_client):
        import openpyxl

        response = admin_client.get('/exports/download-template/device')
        assert response.status_code == 200
        workbook = openpyxl.load_workbook(io.BytesIO(response.data))
        sheet = workbook.active
        assert [c.value for c in sheet[1]] == IMPORT_TEMPLATES['device']['headers']
        assert sheet.freeze_panes == 'C2'
        by_field = {field: sheet.cell(2, col) for col, (_label, field) in
                    enumerate(IMPORT_TEMPLATES['device']['fields'], 1)}
        for field, cell in by_field.items():
            validations = [v for v in sheet.data_validations.dataValidation
                           if cell.coordinate in v.sqref]
            assert len(validations) == 1, field
            validation = validations[0]
            assert validation.showInputMessage is True
            assert validation.prompt and len(validation.prompt) <= 255
            assert f'{cell.column_letter}5000' in validation.sqref
            assert cell.fill.fgColor.rgb == '00FFF2CC'
        for field in ('build_date', 'license_start', 'license_expiry'):
            cell = by_field[field]
            validation = next(v for v in sheet.data_validations.dataValidation
                              if cell.coordinate in v.sqref)
            assert validation.type == 'date'
            assert validation.showErrorMessage is True
            assert validation.allow_blank is True
            assert 'YYYY-MM-DD' in validation.prompt
            assert isinstance(cell.value, (date, datetime))
            assert cell.number_format == 'yyyy-mm-dd'
            assert sheet[f'{cell.column_letter}5000'].number_format == 'yyyy-mm-dd'
        for field in ('serial_number', 'model', 'username', 'rule_version', 'os_version'):
            assert by_field[field].number_format == '@'
        for field in ('port', 'rack_start_u', 'rack_occupy_u', 'rated_power_w', 'device_id'):
            validation = next(v for v in sheet.data_validations.dataValidation
                              if by_field[field].coordinate in v.sqref)
            assert validation.type == 'whole'
            assert validation.showErrorMessage is True
        guide = workbook['填写说明']
        guide_text = '\n'.join(str(cell.value or '') for row in guide for cell in row)
        assert 'YYYY-MM-DD' in guide_text
        assert '不附加00:00:00' in guide_text
        assert '不是Excel行号' in guide_text
        assert '粘贴数据可能绕过Excel校验' in guide_text
        from utils.import_template_guidance import NETWORK_INTERFACE_EXAMPLE, MEETING_INTERFACE_EXAMPLE
        assert NETWORK_INTERFACE_EXAMPLE in guide_text
        assert MEETING_INTERFACE_EXAMPLE in guide_text
        assert '自动' in guide_text
        for field in ('rack_start_u', 'rack_occupy_u'):
            validation = next(v for v in sheet.data_validations.dataValidation
                              if by_field[field].coordinate in v.sqref)
            assert '设备占用23-24U，起始U位填23，占用U数填2' in validation.prompt
            assert '不是结束U号' in validation.prompt
        assert by_field['rack_start_u'].value == 23
        assert by_field['rack_occupy_u'].value == 2
        assert '设备占用23-24U，起始U位填23，占用U数填2' in guide_text
        workbook.close()

    def test_formatted_device_template_roundtrips_dates_and_text(self, admin_client, app):
        import openpyxl

        with app.app_context():
            db.session.add(Customer(name='格式验证客户'))
            db.session.commit()
        response = admin_client.get('/exports/download-template/device')
        workbook = openpyxl.load_workbook(io.BytesIO(response.data))
        sheet = workbook.active
        values = {
            'customer_name': '格式验证客户', 'device_name': '格式验证设备',
            'serial_number': '001234567890123456', 'rule_version': '2026-08-07',
            'build_date': date(2024, 10, 10), 'license_start': date(2024, 10, 10),
            'license_expiry': date(2027, 9, 30),
        }
        for col, (_header, field) in enumerate(IMPORT_TEMPLATES['device']['fields'], 1):
            sheet.cell(2, col).value = values.get(field)
        body = io.BytesIO()
        workbook.save(body)
        workbook.close()
        body.seek(0)
        result = admin_client.post('/api/v2/devices/import', data={
            'import_file': (body, 'formatted.xlsx'), 'mode': 'create',
        }, content_type='multipart/form-data')
        assert result.status_code == 200, result.get_json()
        assert result.get_json()['data']['create'] == 1
        assert result.get_json()['data']['failed'] == 0
        with app.app_context():
            device = Device.query.filter_by(device_name='格式验证设备').one()
            for field, value in values.items():
                if field != 'customer_name':
                    assert getattr(device, field) == value, field

    def test_every_batch_import_has_downloadable_template(self, admin_client):
        """全部现有批量导入模块都能下载同口径 xlsx 模板。"""
        import openpyxl

        for module, definition in IMPORT_TEMPLATES.items():
            assert len(definition['headers']) == len(definition['example']), module
            assert len(definition['headers']) == len(set(definition['headers'])), module
            assert len(definition['fields']) == len({field for _header, field in definition['fields']}), module
            response = admin_client.get(f'/exports/download-template/{module}')
            assert response.status_code == 200, module
            assert response.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            workbook = openpyxl.load_workbook(io.BytesIO(response.data), read_only=True)
            sheet = workbook.active
            headers = [cell.value for cell in sheet[1]]
            assert headers == definition['headers'], module
            assert any(cell.value not in (None, '') for cell in sheet[2]), module

    def test_template_permission_matches_import_permission(self, op_client, sales_client,
                                                           viewer_client):
        assert op_client.get('/exports/download-template/device').status_code == 200
        assert op_client.get('/exports/download-template/inspection').status_code == 200
        assert op_client.get('/exports/download-template/fault').status_code == 200
        assert op_client.get('/exports/download-template/spare').status_code == 200
        assert op_client.get('/exports/download-template/stock').status_code == 200
        assert sales_client.get('/exports/download-template/customer').status_code == 200
        assert sales_client.get('/exports/download-template/device').status_code == 403
        assert viewer_client.get('/exports/download-template/device').status_code == 403

    def test_device_template_login_method_dictionary_contains_local_serial(self, admin_client):
        """设备编辑字典和导入模板必须共用登录方式选项。"""
        import openpyxl

        response = admin_client.get('/exports/download-template/device')
        assert response.status_code == 200
        workbook = openpyxl.load_workbook(io.BytesIO(response.data), read_only=True)
        dictionary = workbook['数据字典']
        headers = [cell.value for cell in dictionary[1]]
        login_method_column = headers.index('登录方式') + 1
        values = [dictionary.cell(row, login_method_column).value
                  for row in range(2, dictionary.max_row + 1)]
        assert '本地串口' in values

    def test_unknown_template_is_404(self, admin_client):
        assert admin_client.get('/exports/download-template/not-exists').status_code == 404

    def test_full_editable_field_contract(self):
        """模板必须覆盖各模块可编辑字段，且不能再出现后端不读取的假字段。"""
        expected = {
            'device': {
                'customer_name', 'device_name', 'device_type', 'brand', 'model',
                'serial_number', 'network_type', 'ip_address', 'port', 'username',
                'password', 'login_method', 'rack_location', 'rack_name', 'location',
                'rack_start_u', 'rack_occupy_u', 'power_supply', 'interface',
                'os_version', 'rule_version', 'build_date', 'license_start',
                'license_expiry', 'is_maintenance', 'is_in_use',
                'remark', 'rated_power_w', 'device_id',
            },
            'customer': {
                'name', 'contact_person', 'phone', 'email', 'region_name', 'city',
                'address', 'category_name', 'level', 'office', 'office_room',
                'map_location', 'has_onsite', 'onsite_contact', 'onsite_phone',
                'onsite_office', 'has_drill', 'inspection_frequency',
                'contract_start_date', 'contract_end_date', 'source', 'remark',
                'parent_name',
            },
            'spare': {
                'code', 'name', 'category', 'brand', 'model', 'specification',
                'unit', 'min_stock', 'reference_price', 'warranty_months',
                'manufacturer', 'serial_number', 'remark',
            },
            'inspection': {
                'customer_name', 'title', 'inspector', 'inspection_date',
                'location', 'overall_status', 'conclusion',
            },
            'fault': {
                'customer_name', 'title', 'handler', 'fault_time', 'fault_type',
                'fault_category', 'fault_description', 'impact_range', 'fault_cause',
                'solution', 'result', 'handling_started_at', 'recovery_time',
            },
            'stock': {'spare_name', 'location', 'quantity', 'unit_price'},
        }
        for module, fields in expected.items():
            assert {field for _header, field in IMPORT_TEMPLATES[module]['fields']} == fields

    def test_device_template_order_matches_device_list(self):
        """客户列用于归属；设备业务列按列表顺序排列，机柜字段不得再被拆散。"""
        assert list(IMPORT_TEMPLATES['device']['fields'][:15]) == [
            ('客户', 'customer_name'),
            ('名称', 'device_name'),
            ('类型', 'device_type'),
            ('机房位置', 'rack_location'),
            ('机柜号', 'rack_name'),
            ('安装位置', 'location'),
            ('起始U位', 'rack_start_u'),
            ('占用U数', 'rack_occupy_u'),
            ('电源配置', 'power_supply'),
            ('额定功率', 'rated_power_w'),
            ('品牌', 'brand'),
            ('型号', 'model'),
            ('序列号', 'serial_number'),
            ('IP', 'ip_address'),
            ('网络类型', 'network_type'),
        ]

    def test_device_template_accepts_legacy_header_aliases(self):
        mapping = get_import_field_mapping('device')
        assert mapping['客户'] == mapping['所属客户'] == 'customer_name'
        assert mapping['名称'] == mapping['设备名称'] == 'device_name'
        assert mapping['类型'] == mapping['设备类型'] == 'device_type'
        assert mapping['IP'] == mapping['IP地址'] == 'ip_address'
        assert mapping['建设时间'] == mapping['建设日期'] == 'build_date'
        assert mapping['授权开始'] == mapping['授权开始日期'] == 'license_start'
        assert mapping['授权截止'] == mapping['授权截止日期'] == 'license_expiry'
        assert mapping['证书到期日期'] == mapping['证书到期'] == 'cert_expiry_date'

    def test_templates_cover_registered_form_profiles(self):
        """新增编辑字段后若忘记同步模板，本测试必须立即失败。"""
        aliases = {
            'device': {'rack_slot': 'rack_start_u'},
            'inspection': {'inspector_name': 'inspector'},
        }
        for module in ('device', 'customer', 'inspection', 'fault', 'spare'):
            template_fields = {
                field for _header, field in IMPORT_TEMPLATES[module]['fields']
            }
            form_fields = {
                aliases.get(module, {}).get(field, field)
                for field in get_entity_schema(module).profiles['form']
            }
            if module == 'device':
                # 证书日期保留详情与旧表导入，但用户已明确从新模板移除。
                form_fields.discard('cert_expiry_date')
            assert form_fields <= template_fields, (
                f'{module} 模板缺少编辑字段：{sorted(form_fields - template_fields)}')


class TestDeviceCustomerImport:
    @pytest.mark.parametrize('date_storage', ['slash_text', 'native_date', 'datetime_text'])
    def test_lifecycle_dates_create_update_and_readback(self, admin_client, app,
                                                      date_storage):
        """用户原表头/斜杠日期经过预检与确认后，DB、列表和详情均必须保留。"""
        fields = ('build_date', 'license_start', 'license_expiry', 'cert_expiry_date')
        headers = ['客户', '名称', '建设时间', '授权开始日期', '授权截止日期', '证书到期日期']
        samples = [
            ['2021/1/1', '2025/3/27', '2026/3/26', '2026/3/26'],
            ['2013/1/1', '', '', ''],
            ['2023/6/18', '2023/6/18', '2026/6/27', '2026/6/27'],
            ['', '', '', ''],
            ['2021/1/1', '2025/7/29', '2026/7/28', '2026/7/28'],
            ['2024/1/1', '2024/1/1', '2027/10/12', '2027/10/12'],
            ['2023/1/1', '', '', ''],
            ['2024/10/10', '2024/10/10', '2027/9/30', '2027/9/30'],
        ]
        customer_name = '日期全链路客户'
        with app.app_context():
            customer = Customer(name=customer_name)
            db.session.add(customer)
            db.session.commit()
            customer_id = customer.id

        def cell_date(value):
            if not value or date_storage == 'slash_text':
                return value
            parsed = datetime.strptime(value, '%Y/%m/%d')
            return parsed if date_storage == 'native_date' else str(parsed)

        def workbook(rows):
            return _xlsx(headers, [
                [customer_name, f'日期设备-{index}', *map(cell_date, values)]
                for index, values in enumerate(rows)
            ])

        def run_import(rows, mode):
            content = workbook(rows)
            preview = admin_client.post('/api/v2/devices/import', data={
                'import_file': (io.BytesIO(content), '日期设备.xlsx'),
                'mode': mode, 'dry_run': '1',
            }, content_type='multipart/form-data')
            assert preview.status_code == 200, preview.get_json()
            result = preview.get_json()['data']
            assert result['committed'] is False
            assert result['failed'] == result['skipped'] == 0, result
            assert result[mode] == len(rows), result
            response = admin_client.post('/api/v2/devices/import', data={
                'import_file': (io.BytesIO(content), '日期设备.xlsx'),
                'mode': mode, 'batch_id': result['batch_id'],
            }, content_type='multipart/form-data')
            assert response.status_code == 200, response.get_json()
            assert response.get_json()['data'][mode] == len(rows)
            assert response.get_json()['data']['committed'] is True

        def assert_readback(rows):
            items = admin_client.get('/api/devices', query_string={
                'customer_id': customer_id,
            }).get_json()['data']['items']
            assert len(items) == len(rows)
            by_name = {item['device_name']: item for item in items}
            for index, values in enumerate(rows):
                item = by_name[f'日期设备-{index}']
                expected = {
                    field: datetime.strptime(value, '%Y/%m/%d').date().isoformat()
                    if value else '' for field, value in zip(fields, values)
                }
                assert {field: item[field] for field in fields} == expected
                detail = admin_client.get(f"/api/devices/{item['id']}").get_json()['data']
                assert {field: detail[field] for field in fields} == expected
                with app.app_context():
                    device = db.session.get(Device, item['id'])
                    assert {field: getattr(device, field).isoformat()
                            if getattr(device, field) else '' for field in fields} == expected

        run_import(samples, 'create')
        assert_readback(samples)
        updates = [['2024/10/10', '2025/3/27', '2027/10/12', '2027/9/30']
                   for _row in samples]
        run_import(updates, 'update')
        assert_readback(updates)

    def test_device_template_imports_all_editable_fields(self, admin_client, app):
        with app.app_context():
            customer = Customer(name='完整设备模板客户')
            db.session.add_all([customer, NetworkType(name='内网', sort_order=1)])
            db.session.commit()
            customer_id = customer.id
        values = {
            'customer_name': '完整设备模板客户', 'device_name': 'FULL-IMPORT-DEVICE',
            'device_type': '交换机', 'brand': '华为', 'model': 'S5735',
            'serial_number': 'SN-FULL', 'network_type': '内网',
            'ip_address': '10.0.0.88', 'port': 2222, 'username': 'ops',
            'password': 'Secret#123', 'login_method': 'SSH',
            'rack_location': '9楼机房', 'rack_name': '4', 'location': '正面',
            'rack_start_u': 27, 'rack_occupy_u': 4, 'power_supply': '双电源',
            'interface': 'GE0/0/1、GE0/0/2', 'os_version': 'V1',
            'rule_version': datetime(2026, 8, 7, 0, 0), 'build_date': date(2026, 1, 2),
            'license_start': datetime(2026, 2, 1, 0, 0), 'license_expiry': '2027/2/1',
            'is_maintenance': '是',
            'is_in_use': '是', 'remark': '完整字段导入',
        }
        data = _xlsx(IMPORT_TEMPLATES['device']['headers'], [
            _template_row('device', values),
        ])
        response = admin_client.post('/api/v2/devices/import', data={
            'import_file': (io.BytesIO(data), 'devices-full.xlsx'),
        }, content_type='multipart/form-data')
        assert response.status_code == 200, response.get_json()
        assert response.get_json()['data']['created'] == 1
        with app.app_context():
            device = Device.query.filter_by(device_name='FULL-IMPORT-DEVICE').one()
            assert device.customer_id == customer_id
            assert device.device_type == '交换机'
            assert device.brand == '华为'
            assert device.model == 'S5735'
            assert device.serial_number == 'SN-FULL'
            assert device.network_type == '内网'
            assert device.ip_address == '10.0.0.88'
            assert device.port == 2222
            assert device.username == 'ops'
            assert decrypt_password(device.password_encrypted) == 'Secret#123'
            assert device.login_method == 'SSH'
            assert device.location == '正面'
            assert device.power_supply == '双电源'
            assert device.interface and parse_json(device.interface) == ['GE0/0/1', 'GE0/0/2']
            assert device.os_version == 'V1'
            assert device.rule_version == '2026-08-07'
            assert device.build_date.isoformat() == '2026-01-02'
            assert device.license_start.isoformat() == '2026-02-01'
            assert device.license_expiry.isoformat() == '2027-02-01'
            assert device.cert_expiry_date is None
            assert device.is_maintenance is True
            assert device.is_in_use is True
            assert device.remark == '完整字段导入'
            rack = Rack.query.filter_by(
                customer_id=customer_id, name='4', location='9楼机房').one()
            install = RackInstall.query.filter_by(device_id=device.id, rack_id=rack.id).one()
            assert (install.start_u, install.occupy_u) == (27, 4)

    def test_device_import_rejects_unknown_network_type(self, admin_client, app):
        with app.app_context():
            db.session.add(Customer(name='未知网络类型客户'))
            db.session.commit()
        data = _xlsx(IMPORT_TEMPLATES['device']['headers'], [
            _template_row('device', {
                'customer_name': '未知网络类型客户', 'device_name': 'BAD-NETWORK-TYPE',
                'network_type': '模板里乱填的类型', 'is_in_use': '是',
            }),
        ])
        response = admin_client.post('/api/v2/devices/import', data={
            'import_file': (io.BytesIO(data), 'bad-network.xlsx'),
        }, content_type='multipart/form-data')
        assert response.status_code == 200
        body = response.get_json()['data']
        assert body['created'] == 0
        assert '未知网络类型' in body['errors'][0]
        assert body['errors_file']['filename'].endswith('.xlsx')

    def test_device_import_preflight_rejects_rack_conflict_atomically(
            self, admin_client, app):
        with app.app_context():
            customer = Customer(name='U位冲突客户')
            db.session.add(customer)
            db.session.flush()
            rack = Rack(customer_id=customer.id, name='1', location='中心机房', total_u=42)
            db.session.add(rack)
            db.session.flush()
            db.session.add(RackInstall(
                rack_id=rack.id, manual_name='原设备', start_u=5, occupy_u=2))
            db.session.commit()
        data = _xlsx(IMPORT_TEMPLATES['device']['headers'], [
            _template_row('device', {
                'customer_name': 'U位冲突客户', 'device_name': '冲突新设备',
                'rack_location': '中心机房', 'rack_name': '1',
                'rack_start_u': 6, 'rack_occupy_u': 2,
            }),
        ])
        response = admin_client.post('/api/v2/devices/import', data={
            'import_file': (io.BytesIO(data), 'rack-conflict.xlsx'),
        }, content_type='multipart/form-data')
        assert response.status_code == 200
        result = response.get_json()['data']
        assert result['failed'] == 1
        assert 'U 位冲突' in result['errors'][0]
        with app.app_context():
            assert Device.query.filter_by(device_name='冲突新设备').count() == 0

    def test_asset_import_accepts_aggregated_u_range(self, admin_client, app):
        with app.app_context():
            customer = Customer(name='聚合U位客户')
            db.session.add(customer)
            db.session.commit()
            customer_id = customer.id
        data = _xlsx(['客户', '名称', '机房位置', '机柜号', '起始U位'], [[
            '聚合U位客户', '聚合U位设备', '中心机房', '2', '27U-30U',
        ]])
        response = admin_client.post('/api/v2/devices/import', data={
            'import_file': (io.BytesIO(data), 'aggregated-u.xlsx'),
        }, content_type='multipart/form-data')
        assert response.status_code == 200, response.get_json()
        with app.app_context():
            device = Device.query.filter_by(
                customer_id=customer_id, device_name='聚合U位设备').one()
            install = RackInstall.query.filter_by(device_id=device.id).one()
            assert (install.start_u, install.occupy_u) == (27, 4)

    def test_customer_template_imports_extended_editable_fields(self, admin_client, app):
        with app.app_context():
            region = Region(name='完整模板地区')
            category = CustomerCategory(name='完整模板单位类别')
            db.session.add_all([region, category])
            db.session.commit()
            region_id, category_id = region.id, category.id
        data = _xlsx(IMPORT_TEMPLATES['customer']['headers'], [
            _template_row('customer', {
                'name': '完整模板客户', 'contact_person': '张三',
                'phone': '13800000000', 'email': 'full@example.com',
                'region_name': '完整模板地区', 'city': '鹰潭市',
                'address': '完整地址', 'category_name': '完整模板单位类别',
                'level': '重点',
                'office': '信息中心', 'office_room': 'A栋301',
                'map_location': '28.27,117.03', 'has_onsite': '是',
                'onsite_contact': '李四', 'onsite_phone': '13900000000',
                'onsite_office': 'B栋201',
                'has_drill': '否', 'inspection_frequency': '每季度',
                'contract_start_date': '2026-01-01',
                'contract_end_date': '2026-12-31', 'source': '批量导入',
                'remark': '完整字段导入',
            }),
        ])
        response = admin_client.post('/api/v2/customers/import', data={
            'import_file': (io.BytesIO(data), 'customers-full.xlsx'),
        }, content_type='multipart/form-data')
        assert response.status_code == 200, response.get_json()
        with app.app_context():
            customer = Customer.query.filter_by(name='完整模板客户').one()
            assert customer.contact_person == '张三'
            assert customer.phone == '13800000000'
            assert customer.email == 'full@example.com'
            assert customer.region_id == region_id
            assert customer.city == '鹰潭市'
            assert customer.address == '完整地址'
            assert customer.category_id == category_id
            assert customer.level == '重点'
            assert customer.office == '信息中心'
            assert customer.office_room == 'A栋301'
            assert customer.map_location == '28.27,117.03'
            assert customer.has_onsite is True
            assert customer.onsite_contact == '李四'
            assert customer.onsite_phone == '13900000000'
            assert customer.onsite_office == 'B栋201'
            assert customer.has_drill is False
            assert customer.inspection_frequency == '每季度'
            assert customer.contract_start_date.isoformat() == '2026-01-01'
            assert customer.contract_end_date.isoformat() == '2026-12-31'
            assert customer.source == '批量导入'
            assert customer.remark == '完整字段导入'


class TestSpareImport:
    def test_import_spare_parts(self, admin_client, app):
        data = _xlsx(IMPORT_TEMPLATES['spare']['headers'], [
            _template_row('spare', {
                'code': 'SP-001', 'name': '光模块SFP', 'category': '光模块',
                'brand': '华为', 'model': 'SFP-GE', 'specification': '千兆',
                'unit': '个', 'min_stock': 5, 'reference_price': 450.5,
                'warranty_months': 12, 'manufacturer': '华为技术',
                'serial_number': 'SFP-SN-1', 'remark': '完整备件字段',
            }),
            _template_row('spare', {
                'code': 'SP-002', 'name': '电源模块', 'category': '电源',
                'specification': '48V', 'unit': '个', 'min_stock': 2,
                'remark': '备用',
            }),
            _template_row('spare', {
                'code': 'SP-001', 'name': '光模块SFP-重复', 'category': '光模块',
            }),
        ])  # 编码重复应跳过
        r = admin_client.post('/api/spare-parts/import',
                              data={'import_file': (io.BytesIO(data), 'spare.xlsx')},
                              content_type='multipart/form-data')
        assert r.status_code == 200, r.get_json()
        d = r.get_json()['data']
        assert d['success'] == 2
        assert d['skipped'] == 1
        with app.app_context():
            assert SparePart.query.filter_by(code='SP-001').count() == 1
            part = SparePart.query.filter_by(code='SP-001').one()
            assert part.brand == '华为'
            assert part.model == 'SFP-GE'
            assert part.category == '光模块'
            assert part.specification == '千兆'
            assert part.unit == '个'
            assert part.min_stock == 5
            assert part.reference_price == 450.5
            assert part.warranty_months == 12
            assert part.manufacturer == '华为技术'
            assert part.serial_number == 'SFP-SN-1'
            assert part.remark == '完整备件字段'

    def test_import_spare_stocks(self, admin_client, app):
        with app.app_context():
            p = SparePart(name='硬盘1T', code='HDD-1T')
            db.session.add(p)
            db.session.commit()
            pid = p.id
        data = _xlsx(
            ['备件名称', '位置', '数量', '单价'],
            [['硬盘1T', 'A柜-01', '10', '450'],
             ['硬盘1T', 'A柜-01', '5', '450'],    # 同名库位累加
             ['不存在的备件', 'B柜', '1', '1']])   # 报错行
        r = admin_client.post('/api/spare-stocks/import',
                              data={'import_file': (io.BytesIO(data), 'stock.xlsx')},
                              content_type='multipart/form-data')
        assert r.status_code == 200, r.get_json()
        d = r.get_json()['data']
        assert d['success'] == 1
        assert len(d['errors']) == 1
        with app.app_context():
            s = SpareStock.query.filter_by(spare_part_id=pid).first()
            assert s.quantity == 15  # 10 + 5 累加


class TestInspectionFaultImport:
    def test_import_inspections(self, admin_client, app):
        with app.app_context():
            c = Customer(name='巡检导入客户')
            db.session.add(c)
            db.session.commit()
            cid = c.id
        data = _xlsx(
            IMPORT_TEMPLATES['inspection']['headers'],
            [_template_row('inspection', {
                'customer_name': '巡检导入客户', 'title': '季度巡检',
                'inspector': '张工', 'inspection_date': '2026-08-01',
                'location': '机房A', 'overall_status': '正常', 'conclusion': '无异常',
            }), _template_row('inspection', {
                'customer_name': '不存在的客户', 'title': '无效巡检',
            })])
        r = admin_client.post('/api/inspections/import',
                              data={'import_file': (io.BytesIO(data), 'insp.xlsx')},
                              content_type='multipart/form-data')
        assert r.status_code == 200, r.get_json()
        d = r.get_json()['data']
        assert d['success'] == 1
        assert len(d['errors']) == 1
        with app.app_context():
            i = Inspection.query.filter_by(customer_id=cid).first()
            assert i.title == '季度巡检'
            assert i.inspector == '张工'
            assert i.inspection_date.isoformat() == '2026-08-01'
            assert i.location == '机房A'
            assert i.overall_status == '正常'
            assert i.conclusion == '无异常'

    def test_import_faults(self, admin_client, app):
        with app.app_context():
            c = Customer(name='故障导入客户')
            db.session.add(c)
            db.session.commit()
            cid = c.id
        data = _xlsx(IMPORT_TEMPLATES['fault']['headers'], [
            _template_row('fault', {
                'customer_name': '故障导入客户', 'title': '交换机离线',
                'handler': '李工', 'fault_time': '2026-08-02 09:30',
                'fault_type': '网络故障', 'fault_category': '网络/链路/中断',
                'fault_description': '端口down', 'impact_range': '业务中断30分钟',
                'fault_cause': '光纤松动', 'solution': '重新插拔',
                'result': '已解决', 'recovery_time': '2026-08-02 10:00',
            }),
        ])
        r = admin_client.post('/api/faults/import',
                              data={'import_file': (io.BytesIO(data), 'fault.xlsx')},
                              content_type='multipart/form-data')
        assert r.status_code == 200, r.get_json()
        d = r.get_json()['data']
        assert d['success'] == 1
        with app.app_context():
            f = Fault.query.filter_by(customer_id=cid).first()
            assert f.title == '交换机离线'
            assert f.result == '已解决'
            assert f.handler == '李工'
            assert f.fault_type == '网络故障'
            assert (f.fault_category_level1, f.fault_category_level2,
                    f.fault_category_level3) == ('网络', '链路', '中断')
            assert f.fault_description == '端口down'
            assert f.impact_range == '业务中断30分钟'
            assert f.fault_cause == '光纤松动'
            assert f.solution == '重新插拔'
            assert f.fault_time.strftime('%Y-%m-%d %H:%M') == '2026-08-02 09:30'
            assert f.recovery_time.strftime('%Y-%m-%d %H:%M') == '2026-08-02 10:00'
