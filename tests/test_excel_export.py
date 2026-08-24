"""Temporary spreadsheet responses must remove generated files."""
import os

from openpyxl import load_workbook

from blueprints.task_schedule import TASK_STATUS_EXCEL_STYLES
from utils.constants import (
    TASK_CANCELLED,
    TASK_CONTRACT_REVIEW,
    TASK_DONE,
    TASK_PENDING,
    TASK_REVIEWING,
    TASK_RUNNING,
)
from utils.excel_export import export_xlsx, send_temp_export


def test_send_temp_export_cleans_file_on_close(app):
    with app.test_request_context('/'):
        path, name = export_xlsx(['列'], [['值']], 'test.xlsx')
        assert os.path.exists(path)
        response = send_temp_export(path, name)
        response.close()
        assert not os.path.exists(path)


def test_task_status_value_styles_match_web_semantic_colors():
    statuses = [
        TASK_PENDING,
        TASK_RUNNING,
        TASK_REVIEWING,
        TASK_DONE,
        TASK_CANCELLED,
        TASK_CONTRACT_REVIEW,
    ]
    expected = {
        TASK_PENDING: ('FAECD8', 'E6A23C'),
        TASK_RUNNING: ('D9ECFF', '409EFF'),
        TASK_REVIEWING: ('E9E9EB', '909399'),
        TASK_DONE: ('E1F3D8', '67C23A'),
        TASK_CANCELLED: ('E9E9EB', '909399'),
        TASK_CONTRACT_REVIEW: ('FDE2E2', 'F56C6C'),
    }
    path, _ = export_xlsx(
        ['完成状态'],
        [[status] for status in statuses],
        'task-status.xlsx',
        column_value_styles={1: TASK_STATUS_EXCEL_STYLES},
    )
    try:
        sheet = load_workbook(path).active
        for row_index, status in enumerate(statuses, 2):
            cell = sheet.cell(row=row_index, column=1)
            fill, font = expected[status]
            assert cell.fill.fill_type == 'solid'
            assert cell.fill.fgColor.rgb.endswith(fill)
            assert cell.font.color.rgb.endswith(font)
            assert cell.font.bold is True
            assert cell.alignment.horizontal == 'center'
    finally:
        os.remove(path)
